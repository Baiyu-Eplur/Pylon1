from __future__ import annotations

import argparse
import csv
import json
from collections import defaultdict
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np

ROOT = Path(__file__).resolve().parents[1]


def configure_fonts() -> None:
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
    ]
    names: list[str] = []
    for path in candidates:
        if path.exists():
            font_manager.fontManager.addfont(str(path))
            names.append(font_manager.FontProperties(fname=str(path)).get_name())
    if names:
        plt.rcParams["font.sans-serif"] = names + plt.rcParams.get("font.sans-serif", [])
    plt.rcParams["axes.unicode_minus"] = False

CRITERIA = {
    "C2_and_C4": "C2&C4",
    "C2_sustained_negative_damping": "C2负阻尼",
    "C4_rolling_response_growth": "C4增长",
    "C6_large_response": "C6大响应",
    "C7_clearance_limit": "C7",
}


def load_cases(manifest_paths: list[Path]) -> dict[str, dict[str, Any]]:
    cases: dict[str, dict[str, Any]] = {}
    for path in manifest_paths:
        manifest = json.loads(path.read_text(encoding="utf-8"))
        for case in manifest["cases"]:
            label = case["case_label"]
            out_dir = Path(case["output_dir"])
            completed = out_dir.exists() and bool(list(out_dir.glob("*_MAX_DISP.csv")))
            cases[label] = {
                "case_label": label,
                "L_m": float(case["L_m"]),
                "Sag_m": float(case["Sag_m"]),
                "u_star": float(case["u_star"]),
                "seed": int(case["seed"]),
                "output_dir": str(out_dir),
                "completed": completed,
                "failed": not completed,
            }
    return cases


def read_audits(audit_root: Path, cases: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for audit_path in sorted(audit_root.glob("*/time_step_coverage_audit.json")):
        label = audit_path.parent.name
        if label not in cases:
            continue
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        row = dict(cases[label])
        for key in CRITERIA:
            row[key] = float(audit["summary"][key]["coverage_fraction"])
        rows.append(row)
    return rows


def mean(values: list[float]) -> float:
    return float(sum(values) / len(values)) if values else 0.0


def summarize(cases: dict[str, dict[str, Any]], audits: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped_cases: dict[tuple[float, float], list[dict[str, Any]]] = defaultdict(list)
    grouped_audits: dict[tuple[float, float], list[dict[str, Any]]] = defaultdict(list)
    for case in cases.values():
        grouped_cases[(case["L_m"], case["u_star"])].append(case)
    for audit in audits:
        grouped_audits[(audit["L_m"], audit["u_star"])].append(audit)

    rows: list[dict[str, Any]] = []
    for key in sorted(grouped_cases):
        l_value, u_star = key
        planned = grouped_cases[key]
        completed = grouped_audits.get(key, [])
        row: dict[str, Any] = {
            "L_m": l_value,
            "Sag_m": planned[0]["Sag_m"],
            "u_star": u_star,
            "planned_n": len(planned),
            "completed_n": len(completed),
            "failed_n": len(planned) - len(completed),
            "failed_fraction": (len(planned) - len(completed)) / len(planned),
        }
        for criterion in CRITERIA:
            row[criterion] = mean([item[criterion] for item in completed])
        rows.append(row)
    return rows


def pct(value: float) -> str:
    return f"{100.0 * value:.2f}%"


def write_csv(rows: list[dict[str, Any]], path: Path) -> None:
    keys: list[str] = []
    for row in rows:
        for key in row:
            if key not in keys:
                keys.append(key)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys)
        writer.writeheader()
        writer.writerows(rows)


def make_table_rows(rows: list[dict[str, Any]]) -> list[list[str]]:
    table: list[list[str]] = []
    for row in sorted(rows, key=lambda item: item["u_star"]):
        table.append(
            [
                f"{row['u_star']:.2f}",
                pct(row["C2_and_C4"]),
                pct(row["C2_sustained_negative_damping"]),
                pct(row["C4_rolling_response_growth"]),
                pct(row["C6_large_response"]),
                pct(row["C7_clearance_limit"]),
            ]
        )
    return table


def save_table_png(
    rows: list[dict[str, Any]],
    path: Path,
    *,
    title: str,
    scale_y: float = 1.15,
) -> None:
    headers = ["u_star", "C2&C4", "C2负阻尼", "C4增长", "C6大响应", "C7"]
    table_rows = make_table_rows(rows)
    height = max(2.8, 0.38 * (len(table_rows) + 2) * scale_y)
    fig, ax = plt.subplots(figsize=(9.6, height), dpi=180)
    ax.axis("off")
    ax.set_title(title, fontsize=13, fontweight="bold", pad=14)
    table = ax.table(
        cellText=table_rows,
        colLabels=headers,
        cellLoc="center",
        loc="center",
        colWidths=[0.13, 0.16, 0.18, 0.16, 0.16, 0.13],
    )
    table.auto_set_font_size(False)
    table.set_fontsize(9.5)
    table.scale(1.0, 1.35)
    for (r, _c), cell in table.get_celld().items():
        cell.set_linewidth(0.35)
        cell.set_edgecolor("#aab3c2")
        if r == 0:
            cell.set_facecolor("#243447")
            cell.set_text_props(color="white", fontweight="bold")
        elif r % 2 == 0:
            cell.set_facecolor("#f4f7fb")
        else:
            cell.set_facecolor("white")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def save_all_tables_png(summary_rows: list[dict[str, Any]], path: Path) -> None:
    l_values = sorted({row["L_m"] for row in summary_rows})
    headers = ["u_star", "C2&C4", "C2负阻尼", "C4增长", "C6大响应", "C7"]
    fig, axes = plt.subplots(len(l_values), 1, figsize=(10.4, 3.2 * len(l_values)), dpi=180)
    if len(l_values) == 1:
        axes = [axes]
    for ax, l_value in zip(axes, l_values):
        rows = [row for row in summary_rows if row["L_m"] == l_value]
        sag = rows[0]["Sag_m"]
        ax.axis("off")
        ax.set_title(f"L = {l_value:.0f} m, H = {sag:.3f} m", fontsize=12, fontweight="bold", pad=10)
        table = ax.table(
            cellText=make_table_rows(rows),
            colLabels=headers,
            cellLoc="center",
            loc="center",
            colWidths=[0.13, 0.16, 0.18, 0.16, 0.16, 0.13],
        )
        table.auto_set_font_size(False)
        table.set_fontsize(9)
        table.scale(1.0, 1.25)
        for (r, _c), cell in table.get_celld().items():
            cell.set_linewidth(0.3)
            cell.set_edgecolor("#aab3c2")
            if r == 0:
                cell.set_facecolor("#243447")
                cell.set_text_props(color="white", fontweight="bold")
            elif r % 2 == 0:
                cell.set_facecolor("#f4f7fb")
    fig.suptitle("Time-domain galloping coverage by geometry", fontsize=14, fontweight="bold", y=0.995)
    fig.tight_layout(rect=[0, 0, 1, 0.985])
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def write_workbook(summary_rows: list[dict[str, Any]], path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception:
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Geometry_Index"
    ws.append(["L_m", "H_m", "H_over_L", "u_star_count"])
    for l_value in sorted({row["L_m"] for row in summary_rows}):
        rows = [row for row in summary_rows if row["L_m"] == l_value]
        sag = rows[0]["Sag_m"]
        ws.append([l_value, sag, sag / l_value, len(rows)])

    headers = ["u_star", "C2&C4", "C2负阻尼", "C4增长", "C6大响应", "C7"]
    for l_value in sorted({row["L_m"] for row in summary_rows}):
        rows = [row for row in summary_rows if row["L_m"] == l_value]
        sag = rows[0]["Sag_m"]
        sheet = wb.create_sheet(f"L{l_value:.0f}_H{sag:.3f}"[:31])
        sheet.append(headers)
        for item in make_table_rows(rows):
            sheet.append(item)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243447")
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(horizontal="center")
        for col in range(1, 7):
            sheet.column_dimensions[chr(64 + col)].width = 16
    wb.save(path)


def save_surface(summary_rows: list[dict[str, Any]], criterion: str, path: Path, title: str) -> None:
    l_values = sorted({row["L_m"] for row in summary_rows})
    u_values = sorted({row["u_star"] for row in summary_rows})
    grid = np.full((len(l_values), len(u_values)), np.nan)
    for row in summary_rows:
        i = l_values.index(row["L_m"])
        j = u_values.index(row["u_star"])
        grid[i, j] = 100.0 * row[criterion]

    fig, ax = plt.subplots(figsize=(8.6, 5.4), dpi=180)
    im = ax.imshow(
        grid,
        origin="lower",
        aspect="auto",
        extent=[min(u_values), max(u_values), min(l_values), max(l_values)],
        cmap="viridis",
        vmin=0,
        vmax=max(1.0, float(np.nanmax(grid))),
    )
    ax.set_title(title, fontsize=13, fontweight="bold")
    ax.set_xlabel("u_star (m/s)")
    ax.set_ylabel("L (m)")
    ax.set_xticks(u_values)
    ax.set_yticks(l_values)
    cbar = fig.colorbar(im, ax=ax)
    cbar.set_label("Coverage (%)")
    for i, l_value in enumerate(l_values):
        for j, u_star in enumerate(u_values):
            value = grid[i, j]
            if np.isfinite(value):
                color = "white" if value > 0.55 * np.nanmax(grid) else "#1b1f24"
                ax.text(u_star, l_value, f"{value:.1f}", ha="center", va="center", fontsize=8, color=color)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def save_failure_surface(summary_rows: list[dict[str, Any]], path: Path) -> None:
    pseudo_rows = [dict(row, failed_criterion=row["failed_fraction"]) for row in summary_rows]
    save_surface(pseudo_rows, "failed_criterion", path, "Failed/non-normal fraction surface")


def main() -> None:
    configure_fonts()
    parser = argparse.ArgumentParser(description="Build L-u_star coverage tables and separate 2D surfaces.")
    parser.add_argument("--audit-root", required=True)
    parser.add_argument("--out-dir", required=True)
    parser.add_argument("--manifest", action="append", required=True)
    args = parser.parse_args()

    out_dir = (ROOT / args.out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest_paths = [(ROOT / item).resolve() for item in args.manifest]
    cases = load_cases(manifest_paths)
    audits = read_audits((ROOT / args.audit_root).resolve(), cases)
    summary_rows = summarize(cases, audits)

    write_csv(summary_rows, out_dir / "coverage_big_table_source.csv")
    reliability_keys = [
        "L_m",
        "Sag_m",
        "u_star",
        "planned_n",
        "completed_n",
        "failed_n",
        "failed_fraction",
    ]
    write_csv([{key: row[key] for key in reliability_keys} for row in summary_rows], out_dir / "geometry_reliability_summary.csv")
    write_workbook(summary_rows, out_dir / "coverage_big_table_by_geometry.xlsx")

    save_all_tables_png(summary_rows, out_dir / "coverage_tables_all_geometries.png")
    for l_value in sorted({row["L_m"] for row in summary_rows}):
        rows = [row for row in summary_rows if row["L_m"] == l_value]
        sag = rows[0]["Sag_m"]
        save_table_png(
            rows,
            out_dir / f"coverage_table_L{l_value:.0f}_H{sag:.3f}.png",
            title=f"L = {l_value:.0f} m, H = {sag:.3f} m",
        )

    surface_names = {
        "C2_and_C4": "C2C4_surface.png",
        "C2_sustained_negative_damping": "C2_negative_damping_surface.png",
        "C4_rolling_response_growth": "C4_growth_surface.png",
        "C6_large_response": "C6_large_response_surface.png",
        "C7_clearance_limit": "C7_surface.png",
    }
    for criterion, filename in surface_names.items():
        save_surface(
            summary_rows,
            criterion,
            out_dir / filename,
            f"{CRITERIA[criterion]} coverage surface",
        )
    save_failure_surface(summary_rows, out_dir / "failed_fraction_surface.png")

    print(f"Wrote result package to {out_dir}")


if __name__ == "__main__":
    main()
