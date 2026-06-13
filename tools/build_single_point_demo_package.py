from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from cable_analyser.geometry import CableGeometry
from cable_analyser.wind_forces import _linear_extrapolate, _load_coefficient_table


def load_config_safely(path: Path) -> dict[str, Any]:
    """Load generated configs that may contain unescaped Windows paths."""
    import yaml

    text = path.read_text(encoding="utf-8")
    try:
        return yaml.safe_load(text)
    except yaml.YAMLError:
        sanitized_lines: list[str] = []
        for line in text.splitlines():
            stripped = line.strip()
            if ": " in line and '"' in stripped and "\\" in stripped:
                key, value = line.split(":", 1)
                value = value.strip().strip('"')
                sanitized_lines.append(f"{key}: {value!r}")
            else:
                sanitized_lines.append(line)
        return yaml.safe_load("\n".join(sanitized_lines))


def load_matrix(path: Path) -> np.ndarray:
    if not path.exists() or path.stat().st_size == 0:
        raise FileNotFoundError(path)
    return np.loadtxt(path)


def ensure_2d(data: np.ndarray) -> np.ndarray:
    return data[np.newaxis, :] if data.ndim == 1 else data


def read_vector(path: Path, n: int | None = None) -> np.ndarray:
    values = np.loadtxt(path)
    values = np.asarray(values, dtype=float)
    if n is not None:
        values = values[:n]
    return values


def observation_nodes(n_nodes: int) -> list[dict[str, Any]]:
    indices = [
        int(round((n_nodes - 1) * 0.25)),
        int(round((n_nodes - 1) * 0.50)),
        int(round((n_nodes - 1) * 0.75)),
    ]
    labels = ["quarter_1", "midspan", "quarter_3"]
    names = ["1/4 span", "Midspan", "3/4 span"]
    return [
        {"label": label, "name": name, "idx": idx, "node_id": idx + 1}
        for label, name, idx in zip(labels, names, indices)
    ]


def adjacent_elements(node_idx: int, n_nodes: int) -> list[int]:
    elements: list[int] = []
    if node_idx > 0:
        elements.append(node_idx)
    if node_idx < n_nodes - 1:
        elements.append(node_idx + 1)
    return elements


def plot_lines(
    path: Path,
    time: np.ndarray,
    series: dict[str, np.ndarray],
    *,
    title: str,
    ylabel: str,
    hline: float | None = None,
) -> None:
    fig, ax = plt.subplots(figsize=(11.2, 5.8), dpi=170)
    colors = ["#2E74B5", "#4A8F5A", "#C7832B"]
    for color, (name, values) in zip(colors, series.items()):
        ax.plot(time, values, lw=1.35, color=color, label=name)
    if hline is not None:
        ax.axhline(hline, color="#B24A4A", lw=1.0, ls="--", label=f"limit = {hline:g}")
        ax.axhline(-hline, color="#B24A4A", lw=1.0, ls="--")
    ax.set_title(title, fontsize=13, weight="bold")
    ax.set_xlabel("Time (s)")
    ax.set_ylabel(ylabel)
    ax.grid(True, color="#D8DEE8", linewidth=0.65)
    ax.legend(loc="best", frameon=True, fontsize=8.5)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def plot_geometry(path: Path, geo: dict[str, np.ndarray], monitors: list[dict[str, Any]]) -> None:
    fig, ax = plt.subplots(figsize=(11.2, 4.6), dpi=170)
    ax.plot(geo["x"], geo["z"], color="#243447", lw=1.6)
    colors = ["#2E74B5", "#4A8F5A", "#C7832B"]
    for color, item in zip(colors, monitors):
        idx = item["idx"]
        ax.scatter([geo["x"][idx]], [geo["z"][idx]], s=56, color=color, zorder=5)
        ax.annotate(
            f"{item['name']}\nNode {item['node_id']}",
            (geo["x"][idx], geo["z"][idx]),
            textcoords="offset points",
            xytext=(0, 13),
            ha="center",
            fontsize=8.5,
            color=color,
            weight="bold",
        )
    ax.set_title("Single-point observation locations", fontsize=13, weight="bold")
    ax.set_xlabel("Span coordinate X (m)")
    ax.set_ylabel("Initial elevation Z (m)")
    ax.grid(True, color="#D8DEE8", linewidth=0.65)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def load_displacements(
    config_path: Path,
    output_dir: Path,
    n_nodes: int,
    max_records: int | None,
) -> tuple[np.ndarray, np.ndarray, int]:
    cfg = load_config_safely(config_path)
    dt = float(cfg["time_history"]["dt"])
    raw = ensure_2d(load_matrix(output_dir / "Dynamic.out"))
    if raw.shape[1] == 3 * n_nodes + 1:
        time = raw[:, 0]
        data = raw[:, 1:]
    else:
        time = np.arange(raw.shape[0], dtype=float) * dt
        data = raw
    if data.shape[1] != 3 * n_nodes:
        raise ValueError(f"Dynamic.out has {data.shape[1]} data columns; expected {3 * n_nodes}")
    raw_rows = int(data.shape[0])
    if max_records is not None and raw_rows > max_records:
        time = time[:max_records]
        data = data[:max_records, :]
    return time, data, raw_rows


def coefficient_series(
    cfg: dict[str, Any],
    monitors: list[dict[str, Any]],
    n_steps: int,
) -> tuple[
    np.ndarray,
    dict[str, np.ndarray],
    dict[str, np.ndarray],
    dict[str, np.ndarray],
    dict[str, np.ndarray],
]:
    wind_dir = ROOT / cfg["paths"]["wind_dir"]
    time_path = ROOT / cfg["paths"]["time_file"]
    coeff_dir = ROOT / cfg["paths"]["aero_coeffs_dir"]

    wind_time = read_vector(time_path, n_steps)
    cd_table = _load_coefficient_table(coeff_dir / "C_D_data.txt")
    cl_table = _load_coefficient_table(coeff_dir / "C_L_data2.txt")
    dcl_per_degree_table = np.loadtxt(coeff_dir / "dC_L.txt")

    cd: dict[str, np.ndarray] = {}
    cl: dict[str, np.ndarray] = {}
    alpha_deg: dict[str, np.ndarray] = {}
    delta_d: dict[str, np.ndarray] = {}
    for item in monitors:
        node_id = item["node_id"]
        key = f"{item['name']} (Node {node_id})"
        u = read_vector(wind_dir / f"NODE_{node_id}_wind_H.txt", n_steps)
        w = read_vector(wind_dir / f"NODE_{node_id}_wind_V.txt", n_steps)
        alpha = np.abs(np.rad2deg(np.arctan2(w, u)))
        cd_values = _linear_extrapolate(alpha, cd_table[:, 0], cd_table[:, 1])
        dcl_per_degree = _linear_extrapolate(alpha, cd_table[:, 0], dcl_per_degree_table)
        dcl_per_radian = dcl_per_degree * 180.0 / math.pi
        cd[key] = cd_values
        cl[key] = _linear_extrapolate(alpha, cl_table[:, 0], cl_table[:, 1])
        alpha_deg[key] = alpha
        delta_d[key] = dcl_per_radian + cd_values
    return wind_time, cd, cl, alpha_deg, delta_d


def plot_den_hartog_curve(path: Path, coeff_dir: Path, params: dict[str, float]) -> dict[str, Any]:
    cd_table = _load_coefficient_table(coeff_dir / "C_D_data.txt")
    dcl_per_degree = np.loadtxt(coeff_dir / "dC_L.txt")
    alpha = cd_table[:, 0]
    cd = cd_table[:, 1]
    dcl_per_radian = dcl_per_degree * 180.0 / math.pi
    delta = cd + dcl_per_radian
    k_per_mps = params["ro_air"] * params["B"] * params["L"] / (4.0 * params["MassM"] * params["omegaN"])
    ucrit = np.full_like(delta, np.nan, dtype=float)
    neg = delta < 0.0
    ucrit[neg] = -params["xi_structural"] / (k_per_mps * delta[neg])

    fig, axes = plt.subplots(2, 1, figsize=(10.8, 7.0), dpi=170, sharex=True)
    axes[0].plot(alpha, delta, color="#2E74B5", lw=1.5, label=r"$\Delta_D=dC_L/d\alpha+C_D$")
    axes[0].axhline(0.0, color="#B24A4A", lw=1.0, ls="--")
    axes[0].fill_between(alpha, delta, 0, where=neg, color="#B24A4A", alpha=0.16, label="Den Hartog unstable")
    axes[0].set_ylabel(r"$\Delta_D$ (per rad)")
    axes[0].set_title("Den Hartog aerodynamic damping term", fontsize=13, weight="bold")
    axes[0].legend(loc="best", fontsize=8.5)
    axes[0].grid(True, color="#D8DEE8", linewidth=0.65)

    axes[1].plot(alpha, ucrit, color="#4A8F5A", lw=1.5)
    axes[1].set_xlabel("Angle of attack |alpha| (deg)")
    axes[1].set_ylabel("U_crit for xi_total=0 (m/s)")
    axes[1].set_ylim(bottom=0)
    axes[1].grid(True, color="#D8DEE8", linewidth=0.65)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)

    return {
        "k_per_mps": float(k_per_mps),
        "delta_min": float(np.min(delta)),
        "delta_negative_fraction_table": float(np.mean(neg)),
        "ucrit_min_mps": float(np.nanmin(ucrit)),
        "ucrit_median_negative_mps": float(np.nanmedian(ucrit[neg])),
    }


def damping_series(
    path: Path,
    monitors: list[dict[str, Any]],
    n_nodes: int,
) -> tuple[np.ndarray, dict[str, np.ndarray], dict[str, np.ndarray]]:
    raw = ensure_2d(load_matrix(path))
    if raw.shape[1] < 4:
        raise ValueError(f"{path} should have at least four columns: time element prev_xi xi_total")
    times = np.unique(np.round(raw[:, 0], 8))
    mean_by_point: dict[str, np.ndarray] = {}
    min_by_point: dict[str, np.ndarray] = {}
    for item in monitors:
        elems = adjacent_elements(item["idx"], n_nodes)
        key = f"{item['name']} (Node {item['node_id']}; Ele {','.join(map(str, elems))})"
        means = np.full(len(times), np.nan, dtype=float)
        mins = np.full(len(times), np.nan, dtype=float)
        for i, t_value in enumerate(times):
            mask = (np.round(raw[:, 0], 8) == t_value) & np.isin(raw[:, 1].astype(int), elems)
            values = raw[mask, 3]
            if len(values):
                means[i] = float(np.mean(values))
                mins[i] = float(np.min(values))
        mean_by_point[key] = means
        min_by_point[key] = mins
    return times, mean_by_point, min_by_point


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, sheets: dict[str, tuple[list[dict[str, Any]], list[str]]]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    header_fill = PatternFill("solid", fgColor="EAF3F8")
    for sheet_name, (rows, fields) in sheets.items():
        ws = wb.create_sheet(sheet_name[:31])
        ws.append(fields)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for row in rows:
            ws.append([row.get(field) for field in fields])
        ws.freeze_panes = "A2"
        for col in ws.columns:
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, max_len + 2), 28)
    wb.save(path)


def build_package(config_path: Path, out_dir: Path) -> dict[str, Any]:
    cfg = load_config_safely(config_path)
    geo = CableGeometry(cfg).generate()
    output_dir = ROOT / cfg["paths"]["output_dir"]
    n_nodes = len(geo["x"])
    monitors = observation_nodes(n_nodes)
    target_records = int(cfg["time_history"]["npt"])
    time, disp, raw_dynamic_rows = load_displacements(config_path, output_dir, n_nodes, target_records)
    n_steps = len(time)

    # In this model the aerodynamic x direction is OpenSees DOF 2; DOF 1 is the span coordinate.
    disp_x = disp[:, 1::3]
    disp_z = disp[:, 2::3]

    out_dir.mkdir(parents=True, exist_ok=True)
    wind_time, cd, cl, alpha_deg, delta_d = coefficient_series(cfg, monitors, n_steps)
    damping_time, xi_mean, xi_min = damping_series(output_dir / "damping_change_log.txt", monitors, n_nodes)

    displacement_rows: list[dict[str, Any]] = []
    disp_fields = ["time_s"]
    for item in monitors:
        prefix = item["label"]
        disp_fields.extend([f"{prefix}_node", f"{prefix}_disp_x_m", f"{prefix}_disp_z_m"])
    for i, t_value in enumerate(time):
        row: dict[str, Any] = {"time_s": float(t_value)}
        for item in monitors:
            prefix = item["label"]
            idx = item["idx"]
            row[f"{prefix}_node"] = item["node_id"]
            row[f"{prefix}_disp_x_m"] = float(disp_x[i, idx])
            row[f"{prefix}_disp_z_m"] = float(disp_z[i, idx])
        displacement_rows.append(row)

    coeff_rows: list[dict[str, Any]] = []
    coeff_fields = ["time_s"]
    for item in monitors:
        prefix = item["label"]
        coeff_fields.extend(
            [f"{prefix}_node", f"{prefix}_alpha_deg", f"{prefix}_C_D", f"{prefix}_C_L", f"{prefix}_delta_D"]
        )
    for i, t_value in enumerate(wind_time):
        row = {"time_s": float(t_value)}
        for item in monitors:
            prefix = item["label"]
            key = f"{item['name']} (Node {item['node_id']})"
            row[f"{prefix}_node"] = item["node_id"]
            row[f"{prefix}_alpha_deg"] = float(alpha_deg[key][i])
            row[f"{prefix}_C_D"] = float(cd[key][i])
            row[f"{prefix}_C_L"] = float(cl[key][i])
            row[f"{prefix}_delta_D"] = float(delta_d[key][i])
        coeff_rows.append(row)

    damping_rows: list[dict[str, Any]] = []
    damping_fields = ["time_s"]
    for item in monitors:
        prefix = item["label"]
        elems = adjacent_elements(item["idx"], n_nodes)
        damping_fields.extend(
            [f"{prefix}_node", f"{prefix}_adjacent_elements", f"{prefix}_xi_mean", f"{prefix}_xi_min"]
        )
    for i, t_value in enumerate(damping_time):
        row = {"time_s": float(t_value)}
        for item in monitors:
            prefix = item["label"]
            elems = adjacent_elements(item["idx"], n_nodes)
            key = f"{item['name']} (Node {item['node_id']}; Ele {','.join(map(str, elems))})"
            row[f"{prefix}_node"] = item["node_id"]
            row[f"{prefix}_adjacent_elements"] = ",".join(map(str, elems))
            row[f"{prefix}_xi_mean"] = None if np.isnan(xi_mean[key][i]) else float(xi_mean[key][i])
            row[f"{prefix}_xi_min"] = None if np.isnan(xi_min[key][i]) else float(xi_min[key][i])
        damping_rows.append(row)

    summary_rows: list[dict[str, Any]] = []
    for item in monitors:
        idx = item["idx"]
        coeff_key = f"{item['name']} (Node {item['node_id']})"
        damping_key = (
            f"{item['name']} (Node {item['node_id']}; "
            f"Ele {','.join(map(str, adjacent_elements(idx, n_nodes)))})"
        )
        xi_values = xi_mean[damping_key]
        xi_valid = xi_values[np.isfinite(xi_values)]
        summary_rows.append(
            {
                "monitor": item["name"],
                "node": item["node_id"],
                "x_over_L": float(geo["x"][idx] / float(cfg["geometry"]["L"])),
                "initial_z_m": float(geo["z"][idx]),
                "max_abs_disp_x_m": float(np.max(np.abs(disp_x[:, idx]))),
                "max_abs_disp_z_m": float(np.max(np.abs(disp_z[:, idx]))),
                "mean_C_D": float(np.mean(cd[coeff_key])),
                "max_C_D": float(np.max(cd[coeff_key])),
                "mean_C_L": float(np.mean(cl[coeff_key])),
                "min_C_L": float(np.min(cl[coeff_key])),
                "mean_delta_D": float(np.mean(delta_d[coeff_key])),
                "min_delta_D": float(np.min(delta_d[coeff_key])),
                "negative_delta_D_fraction_from_wind_angle": float(np.mean(delta_d[coeff_key] < 0.0)),
                "mean_xi": float(np.mean(xi_valid)) if len(xi_valid) else None,
                "min_xi": float(np.min(xi_valid)) if len(xi_valid) else None,
                "negative_xi_fraction_recorded": float(np.mean(xi_valid < 0.0)) if len(xi_valid) else None,
            }
        )

    write_csv(out_dir / "single_point_displacement_timeseries.csv", displacement_rows, disp_fields)
    write_csv(out_dir / "single_point_aero_coeff_timeseries.csv", coeff_rows, coeff_fields)
    write_csv(out_dir / "single_point_damping_timeseries.csv", damping_rows, damping_fields)
    summary_fields = list(summary_rows[0].keys())
    write_csv(out_dir / "single_point_summary.csv", summary_rows, summary_fields)
    write_xlsx(
        out_dir / "single_point_demo_tables.xlsx",
        {
            "summary": (summary_rows, summary_fields),
            "displacement": (displacement_rows, disp_fields),
            "aero_coeff": (coeff_rows, coeff_fields),
            "damping": (damping_rows, damping_fields),
        },
    )

    disp_x_series = {
        f"{item['name']} (Node {item['node_id']})": disp_x[:, item["idx"]]
        for item in monitors
    }
    disp_z_series = {
        f"{item['name']} (Node {item['node_id']})": disp_z[:, item["idx"]]
        for item in monitors
    }
    plot_geometry(out_dir / "monitor_points_geometry.png", geo, monitors)
    plot_lines(
        out_dir / "displacement_x_timeseries.png",
        time,
        disp_x_series,
        title="Horizontal/aerodynamic-x displacement response at observation points",
        ylabel="Displacement x (m)",
    )
    plot_lines(
        out_dir / "displacement_z_timeseries.png",
        time,
        disp_z_series,
        title="Vertical z displacement response at observation points",
        ylabel="Displacement z (m)",
    )
    plot_lines(
        out_dir / "cd_timeseries.png",
        wind_time,
        cd,
        title="Drag coefficient C_D at observation points",
        ylabel="C_D",
    )
    plot_lines(
        out_dir / "cl_timeseries.png",
        wind_time,
        cl,
        title="Lift coefficient C_L at observation points",
        ylabel="C_L",
    )
    plot_lines(
        out_dir / "damping_timeseries.png",
        damping_time,
        xi_mean,
        title="Recorded effective damping ratio near observation points",
        ylabel="Mean adjacent-element xi_total",
        hline=0.0,
    )
    plot_lines(
        out_dir / "den_hartog_delta_timeseries.png",
        wind_time,
        delta_d,
        title="Den Hartog term from local wind angle at observation points",
        ylabel="dC_L/dalpha + C_D (per rad)",
        hline=0.0,
    )
    damping_params = {}
    params_path = ROOT / "inputs_aerodynamic_damping.tcl"
    if params_path.exists():
        for line in params_path.read_text(encoding="utf-8").splitlines():
            parts = line.split()
            if len(parts) == 3 and parts[0] == "set":
                try:
                    damping_params[parts[1]] = float(parts[2])
                except ValueError:
                    pass
    den_hartog_audit = plot_den_hartog_curve(
        out_dir / "den_hartog_delta_curve.png",
        ROOT / cfg["paths"]["aero_coeffs_dir"],
        {
            "MassM": damping_params.get("MassM", float("nan")),
            "B": damping_params.get("B", float(cfg["material"]["Dia"])),
            "L": damping_params.get("L", float(cfg["geometry"]["discretisation"])),
            "ro_air": damping_params.get("ro_air", float(cfg["time_history"]["ro_air"])),
            "xi_structural": damping_params.get("xi_structural", float(cfg["damping"]["xi"])),
            "omegaN": damping_params.get("omegaN", float("nan")),
        },
    )

    meta = {
        "config_path": str(config_path.relative_to(ROOT)),
        "output_dir": str(output_dir.relative_to(ROOT)),
        "package_dir": str(out_dir.relative_to(ROOT)),
        "case": {
            "L_m": float(cfg["geometry"]["L"]),
            "Sag_m": float(cfg["geometry"]["Sag"]),
            "u_star_mps": float(cfg["wind_generation"]["u_star"]),
            "seed": cfg["wind_generation"].get("seed"),
            "dt_s": float(cfg["time_history"]["dt"]),
            "npt": int(cfg["time_history"]["npt"]),
            "duration_s": float(n_steps * float(cfg["time_history"]["dt"])),
            "raw_dynamic_rows": raw_dynamic_rows,
            "records_used": n_steps,
        },
        "direction_note": (
            "The plotted x displacement is the aerodynamic horizontal direction used by the wind "
            "and damping routines, corresponding to OpenSees DOF 2 in this span-aligned model. "
            "OpenSees DOF 1 is the span coordinate."
        ),
        "monitors": [
            {
                "label": item["label"],
                "name": item["name"],
                "node_id": item["node_id"],
                "x_m": float(geo["x"][item["idx"]]),
                "x_over_L": float(geo["x"][item["idx"]] / float(cfg["geometry"]["L"])),
                "initial_z_m": float(geo["z"][item["idx"]]),
                "adjacent_elements": adjacent_elements(item["idx"], n_nodes),
            }
            for item in monitors
        ],
        "outputs": {
            "tables_xlsx": "single_point_demo_tables.xlsx",
            "summary_csv": "single_point_summary.csv",
            "displacement_csv": "single_point_displacement_timeseries.csv",
            "aero_coeff_csv": "single_point_aero_coeff_timeseries.csv",
            "damping_csv": "single_point_damping_timeseries.csv",
            "pngs": [
                "monitor_points_geometry.png",
                "displacement_x_timeseries.png",
                "displacement_z_timeseries.png",
                "cd_timeseries.png",
                "cl_timeseries.png",
                "damping_timeseries.png",
                "den_hartog_delta_timeseries.png",
                "den_hartog_delta_curve.png",
            ],
        },
        "den_hartog_audit": den_hartog_audit,
    }
    (out_dir / "single_point_demo_metadata.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    return meta


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a single-point galloping demonstration package.")
    parser.add_argument(
        "--config",
        default="output/sweeps/fixedL322_broad_ustar_n4096_s2_cont/configs/L322P8_U0P600_SEED20260909.yaml",
        help="Config path for the already executed example case.",
    )
    parser.add_argument(
        "--out-dir",
        default="output/single_point_demo/L322P8_U0P600_SEED20260909",
        help="Output directory for tables, figures, and metadata.",
    )
    args = parser.parse_args()
    meta = build_package(ROOT / args.config, ROOT / args.out_dir)
    print(json.dumps(meta, indent=2))


if __name__ == "__main__":
    main()
