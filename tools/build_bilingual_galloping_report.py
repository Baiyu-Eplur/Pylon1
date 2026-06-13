from __future__ import annotations

import csv
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULT_DIR = ROOT / "output" / "time_step_coverage_audit" / "fixedL322_broad_ustar_n4096_s2"
CSV_PATH = RESULT_DIR / "ustar_coverage_summary.csv"
FAILED_CSV = RESULT_DIR / "failed_cases.csv"
REPORT_DIR = ROOT / "output" / "reports" / "galloping_fixedL322_bilingual"
DOCX_PATH = REPORT_DIR / "Galloping_FixedL322_TimeDomain_Coverage_Report_Bilingual.docx"
CHART_PATH = REPORT_DIR / "core_criteria_uncertainty.png"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(85, 85, 85)
HEADER_FILL = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
BORDER = "B7C9DD"


def load_rows() -> list[dict[str, str]]:
    with CSV_PATH.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def load_failed() -> list[dict[str, str]]:
    with FAILED_CSV.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def pct(value: str | float) -> str:
    return f"{float(value) * 100:.2f}%"


def ci(row: dict[str, str], prefix: str) -> str:
    return f"{pct(row[f'{prefix}__ci95_low_coverage'])}-{pct(row[f'{prefix}__ci95_high_coverage'])}"


def find_font(candidates: list[str], size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def draw_chart(rows: list[dict[str, str]]) -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    width, height = 1800, 1050
    margin_l, margin_r, margin_t, margin_b = 150, 80, 90, 160
    plot_w = width - margin_l - margin_r
    plot_h = height - margin_t - margin_b
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font = find_font(
        [
            r"C:\Windows\Fonts\arial.ttf",
            r"C:\Windows\Fonts\calibri.ttf",
            r"C:\Windows\Fonts\simhei.ttf",
        ],
        32,
    )
    small = find_font([r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\calibri.ttf"], 24)
    title_font = find_font([r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\calibrib.ttf"], 38)

    criteria = [
        ("C2_and_C4", "C2&C4 strict marker", (56, 95, 150)),
        ("C4_rolling_response_growth", "C4 response growth", (214, 91, 91)),
        ("C6_large_response", "C6 large response", (74, 144, 91)),
    ]
    x_vals = [float(row["u_star"]) for row in rows]
    x_min, x_max = min(x_vals), max(x_vals)
    y_max = 1.0

    def sx(x: float) -> float:
        return margin_l + (x - x_min) / (x_max - x_min) * plot_w

    def sy(y: float) -> float:
        return margin_t + (1 - y / y_max) * plot_h

    draw.text((margin_l, 28), "Core galloping coverage curves with approximate 95% CI", fill=(25, 45, 70), font=title_font)
    draw.line((margin_l, margin_t + plot_h, margin_l + plot_w, margin_t + plot_h), fill=(60, 60, 60), width=3)
    draw.line((margin_l, margin_t, margin_l, margin_t + plot_h), fill=(60, 60, 60), width=3)

    for i in range(0, 11):
        y = i / 10
        py = sy(y)
        draw.line((margin_l, py, margin_l + plot_w, py), fill=(228, 232, 238), width=1)
        draw.text((35, py - 14), f"{int(y*100)}%", fill=(80, 80, 80), font=small)

    for row in rows:
        x = float(row["u_star"])
        px = sx(x)
        draw.line((px, margin_t + plot_h, px, margin_t + plot_h + 10), fill=(60, 60, 60), width=2)
        draw.text((px - 24, margin_t + plot_h + 18), f"{x:.1f}", fill=(60, 60, 60), font=small)

    for key, label, color in criteria:
        points = []
        for row in rows:
            x = float(row["u_star"])
            mean = float(row[f"{key}__mean_coverage"])
            lo = float(row[f"{key}__ci95_low_coverage"])
            hi = float(row[f"{key}__ci95_high_coverage"])
            px, py = sx(x), sy(mean)
            points.append((px, py))
            draw.line((px, sy(hi), px, sy(lo)), fill=color, width=3)
            draw.line((px - 10, sy(hi), px + 10, sy(hi)), fill=color, width=3)
            draw.line((px - 10, sy(lo), px + 10, sy(lo)), fill=color, width=3)
            draw.ellipse((px - 8, py - 8, px + 8, py + 8), fill=color, outline=color)
        for a, b in zip(points, points[1:]):
            draw.line((a[0], a[1], b[0], b[1]), fill=color, width=5)

    legend_x = margin_l + 30
    legend_y = margin_t + 28
    for idx, (_, label, color) in enumerate(criteria):
        y = legend_y + idx * 42
        draw.line((legend_x, y + 14, legend_x + 45, y + 14), fill=color, width=5)
        draw.ellipse((legend_x + 18, y + 6, legend_x + 34, y + 22), fill=color)
        draw.text((legend_x + 60, y), label, fill=(35, 35, 35), font=small)

    draw.text((width // 2 - 115, height - 58), "u_star (m/s)", fill=(60, 60, 60), font=font)
    img = img.resize((1200, 700), Image.Resampling.LANCZOS)
    img.save(CHART_PATH)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None, size: float = 9.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    set_run_font(run)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run_font(run, east_asia: str = "Microsoft YaHei") -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def set_table_borders(table, color: str = BORDER) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = "Galloping Time-Domain Coverage Study | 固定跨距时域覆盖率研究"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED
        set_run_font(run)


def add_para(doc: Document, zh: str, en: str | None = None, style: str | None = None, bold_label: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    if bold_label:
        r = p.add_run(bold_label)
        r.bold = True
        set_run_font(r)
    r = p.add_run(zh)
    set_run_font(r)
    if en:
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.35)
        p2.paragraph_format.space_after = Pt(8)
        r2 = p2.add_run(en)
        r2.italic = True
        r2.font.color.rgb = MUTED
        set_run_font(r2)


def add_bullet(doc: Document, zh: str, en: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(zh)
    set_run_font(r)
    p2 = doc.add_paragraph(style="List Bullet")
    p2.paragraph_format.left_indent = Cm(0.9)
    r2 = p2.add_run(en)
    r2.italic = True
    r2.font.color.rgb = MUTED
    set_run_font(r2)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    set_run_font(r)


def add_summary_table(doc: Document, rows: list[dict[str, str]]) -> None:
    headers = ["u*", "完成\nCompleted", "失败\nFailed", "C2&C4", "C2", "C4", "C6"]
    keys = [
        "u_star",
        "completed_n",
        "failed_n",
        "C2_and_C4__mean_coverage",
        "C2_sustained_negative_damping__mean_coverage",
        "C4_rolling_response_growth__mean_coverage",
        "C6_large_response__mean_coverage",
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [950, 1150, 1000, 1200, 1200, 1200, 1200]
    for i, h in enumerate(headers):
        set_cell_width(table.rows[0].cells[i], widths[i])
        set_cell_shading(table.rows[0].cells[i], HEADER_FILL)
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=DARK_BLUE, size=8.8)
    for row in rows:
        cells = table.add_row().cells
        values = []
        for key in keys:
            if key.endswith("coverage"):
                values.append(pct(row[key]))
            elif key == "u_star":
                values.append(f"{float(row[key]):.2f}")
            else:
                values.append(str(row[key]))
        for i, value in enumerate(values):
            set_cell_width(cells[i], widths[i])
            set_cell_text(cells[i], value, size=8.8)
    set_table_borders(table)


def add_uncertainty_table(doc: Document, rows: list[dict[str, str]]) -> None:
    headers = ["u*", "n", "C2&C4 mean", "C2&C4 95% CI", "C4 mean", "C4 95% CI", "C6 mean", "C6 95% CI"]
    widths = [700, 650, 1150, 1500, 1150, 1500, 1150, 1500]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        set_cell_width(table.rows[0].cells[i], widths[i])
        set_cell_shading(table.rows[0].cells[i], HEADER_FILL)
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=DARK_BLUE, size=8.2)
    for row in rows:
        values = [
            f"{float(row['u_star']):.2f}",
            row["completed_n"],
            pct(row["C2_and_C4__mean_coverage"]),
            ci(row, "C2_and_C4"),
            pct(row["C4_rolling_response_growth__mean_coverage"]),
            ci(row, "C4_rolling_response_growth"),
            pct(row["C6_large_response__mean_coverage"]),
            ci(row, "C6_large_response"),
        ]
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_width(cells[i], widths[i])
            set_cell_text(cells[i], value, size=7.7)
    set_table_borders(table)


def add_failed_table(doc: Document, rows: list[dict[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["Case", "u*", "Seed", "Note"]
    widths = [4700, 900, 1300, 2200]
    for i, h in enumerate(headers):
        set_cell_width(table.rows[0].cells[i], widths[i])
        set_cell_shading(table.rows[0].cells[i], HEADER_FILL)
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=DARK_BLUE, size=8.3)
    for row in rows:
        cells = table.add_row().cells
        values = [row["case_label"], f"{float(row['u_star']):.2f}", row["seed"], row["note"]]
        for i, value in enumerate(values):
            set_cell_width(cells[i], widths[i])
            set_cell_text(cells[i], value, size=7.5 if i == 0 else 8.0)
            if i == 0:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_table_borders(table)


def add_source_list(doc: Document) -> None:
    sources = [
        "output/time_step_coverage_audit/fixedL322_broad_ustar_n4096_s2/ustar_coverage_summary.csv",
        "output/time_step_coverage_audit/fixedL322_broad_ustar_n4096_s2/case_coverage.csv",
        "output/time_step_coverage_audit/fixedL322_broad_ustar_n4096_s2/failed_cases.csv",
        "tools/time_step_coverage_audit.py",
        "tools/aggregate_time_step_coverage.py",
        "PROJECT_MEMORY.md, RESEARCH_PLAN.md, GALLOPING_CRITERIA.md",
    ]
    for item in sources:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        r.font.size = Pt(9)
        set_run_font(r)


def build_docx() -> None:
    rows = load_rows()
    failed = load_failed()
    draw_chart(rows)

    doc = Document()
    style_doc(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    r = title.add_run("架空线缆 Galloping 时域覆盖率研究报告")
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(11, 37, 69)
    r.bold = True
    set_run_font(r)
    sub = doc.add_paragraph()
    r = sub.add_run("Bilingual Report on Time-Domain Galloping Coverage for Fixed Span L = 322.8 m")
    r.font.size = Pt(12)
    r.font.color.rgb = MUTED
    r.italic = True
    set_run_font(r)
    meta = doc.add_paragraph()
    r = meta.add_run("Date: 2026-06-03 | Geometry: L = 322.8 m, H = 10.48 m | dt = 0.05 s, npt = 4096")
    r.font.size = Pt(9.5)
    r.font.color.rgb = MUTED
    set_run_font(r)

    add_para(
        doc,
        "本报告整理固定跨距线缆在不同摩擦速度 u_star 下的 galloping 时域覆盖率结果。报告重点展示 C2、C4、C6 与 C2&C4 判据的含义、覆盖率趋势、不确定性带，以及非正常/失败样本的研究解释。",
        "This report summarizes the fixed-span galloping time-domain coverage results for different friction velocities u_star. It focuses on the meaning of C2, C4, C6, and C2&C4, the coverage trends, uncertainty bands, and the interpretation of failed or non-normal runs.",
    )

    doc.add_heading("1. 研究设置 / Study Setup", level=1)
    add_para(
        doc,
        "当前报告对应固定几何基准：跨距 L = 322.8 m，下垂高度 H = 10.48 m，目标动力记录长度为 4096 步，时间步长为 0.05 s，因此标准比较窗口为 204.8 s。",
        "The current report uses the fixed baseline geometry: span L = 322.8 m, sag H = 10.48 m, 4096 target dynamic records, and a time step of 0.05 s. The standard comparison window is therefore 204.8 s.",
    )
    add_para(
        doc,
        "所有完成样本进入覆盖率统计；OpenSees 返回非零或记录不完整的样本单独计入 failed/non-normal 结果，不混入覆盖率均值。这一点很重要，因为失败本身可能反映强非线性响应或数值收敛困难。",
        "Completed samples are used for coverage statistics. OpenSees non-zero exits or incomplete records are stored separately as failed/non-normal results and are not averaged into the coverage values. This is important because failure itself may indicate severe nonlinear response or numerical convergence difficulty.",
    )

    doc.add_heading("2. 判据定义 / Criteria Definitions", level=1)
    add_para(
        doc,
        "C2：负有效阻尼判据。它检查同一时间步内是否有足够比例的单元出现负气动阻尼，并且最小阻尼低于阈值。C2 是机理触发指标，表示风场可能正在向振动系统输入能量。",
        "C2: sustained negative effective damping. It checks whether a sufficient fraction of elements has negative aerodynamic damping at the same time step and whether the minimum damping is below the threshold. C2 is a mechanism indicator showing that the wind may be injecting energy into the vibration system.",
    )
    add_para(
        doc,
        "C4：滚动响应增长判据。它基于 20 s 滚动窗口的全局位移 p95 包络，判断当前响应是否相对前一窗口和初始基准窗口发生显著增长。C4 是动态增长过程指标。",
        "C4: rolling response-growth criterion. It uses the global displacement p95 envelope in a 20 s rolling window and checks whether the current response has grown sufficiently relative to both the previous window and the baseline window. C4 is a dynamic growth-process indicator.",
    )
    add_para(
        doc,
        "C6：大响应幅值判据。它检查全局位移是否超过 10% sag。对于当前 H = 10.48 m，C6 位移阈值为约 1.048 m。C6 是工程后果和严重程度指标。",
        "C6: large-response criterion. It checks whether the global displacement exceeds 10% of sag. For H = 10.48 m, the C6 displacement limit is approximately 1.048 m. C6 is an engineering consequence and severity indicator.",
    )
    add_para(
        doc,
        "C2&C4：严格联合判据。它要求 C2 和 C4 在同一时间步同时成立，因此覆盖率通常较低，但物理解释更严格。",
        "C2&C4: strict combined criterion. It requires C2 and C4 to be true at the same time step, so its coverage is usually low but its physical interpretation is stricter.",
    )

    doc.add_heading("3. 覆盖率均值 / Mean Coverage", level=1)
    add_para(
        doc,
        "表 1 汇总了各 u_star 下的完成样本数量、失败样本数量和主要判据平均覆盖率。覆盖率定义为满足某判据的时间步数量除以总时间步数量。",
        "Table 1 summarizes the number of completed samples, failed samples, and mean coverage for the main criteria at each u_star. Coverage is defined as the number of time steps satisfying a criterion divided by the total number of time steps.",
    )
    add_summary_table(doc, rows)
    add_caption(doc, "表 1 / Table 1. Mean time-step coverage by u_star.")

    doc.add_heading("4. 不确定性带 / Uncertainty Bands", level=1)
    add_para(
        doc,
        "表 2 给出 C2&C4、C4 和 C6 的平均覆盖率以及近似 95% 置信区间。由于部分 u_star 的样本数仍然较小，这些区间应理解为描述性不确定性带，而非最终统计置信结论。",
        "Table 2 reports the mean coverage and approximate 95% confidence intervals for C2&C4, C4, and C6. Because some u_star values still have small sample sizes, these intervals should be interpreted as descriptive uncertainty bands rather than final statistical confidence statements.",
    )
    add_uncertainty_table(doc, rows)
    add_caption(doc, "表 2 / Table 2. Core-criterion mean coverage and approximate 95% confidence intervals.")

    doc.add_paragraph()
    doc.add_picture(str(CHART_PATH), width=Inches(6.25))
    add_caption(doc, "图 1 / Figure 1. C2&C4, C4, and C6 coverage curves with approximate 95% confidence intervals.")

    doc.add_heading("5. 失败样本解释 / Failed or Non-Normal Runs", level=1)
    add_para(
        doc,
        "失败样本不是简单的无效样本。当前检查表明，它们通常已经生成 Dynamic.out、Accel.out、Velocity.out 和 damping_change_log.txt 等文件，说明模型已进入动力求解，只是在强非线性阶段以非零状态退出。",
        "Failed samples are not simply invalid samples. The current inspection shows that they usually generated Dynamic.out, Accel.out, Velocity.out, and damping_change_log.txt, meaning the model entered dynamic analysis but exited with a non-zero status during a severe nonlinear stage.",
    )
    add_failed_table(doc, failed)
    add_caption(doc, "表 3 / Table 3. Failed or non-normal samples retained as auxiliary instability evidence.")

    doc.add_heading("6. 主要结论 / Main Conclusions", level=1)
    add_bullet(
        doc,
        "C6 在 u_star = 0.50 到 1.00 区间呈现最清晰的上升趋势，是当前最稳定的大响应严重程度指标。",
        "C6 shows the clearest increasing trend from u_star = 0.50 to 1.00 and is currently the most stable large-response severity metric.",
    )
    add_bullet(
        doc,
        "C4 在中高风速段保持较高覆盖率，但随机性明显，因此应始终和不确定性带一起报告。",
        "C4 remains elevated in the mid-to-high wind range, but it is seed-sensitive and should always be reported with uncertainty bands.",
    )
    add_bullet(
        doc,
        "C2&C4 很保守，覆盖率较低，适合作为严格同步 galloping 机制标记，而不适合作为唯一概率曲线。",
        "C2&C4 is conservative with low coverage. It is suitable as a strict simultaneous galloping mechanism marker, not as the only probability curve.",
    )
    add_bullet(
        doc,
        "失败/非正常比例应作为辅助曲线保留，因为它可能反映强响应、几何非线性和负阻尼共同导致的求解困难。",
        "The failed/non-normal fraction should be retained as an auxiliary curve because it may reflect solver difficulty caused by severe response, geometric nonlinearity, and negative damping.",
    )

    doc.add_heading("7. 下一步 / Next Step", level=1)
    add_para(
        doc,
        "建议下一步扩展到多个跨距：L = 300, 322.8, 350, 400 m；u_star = 0.40, 0.50, 0.60, 0.80, 1.00；每个点先采用 3 个 seed。所有结果应同时输出 C4、C6、C2&C4 覆盖率、不确定性带和失败率。",
        "The next step should expand to multiple spans: L = 300, 322.8, 350, and 400 m; u_star = 0.40, 0.50, 0.60, 0.80, and 1.00; initially with three seeds per point. Each result should report C4, C6, C2&C4 coverage, uncertainty bands, and the failed-run fraction.",
    )

    doc.add_heading("8. 数据来源 / Data Sources", level=1)
    add_source_list(doc)

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_docx()
    print(DOCX_PATH)
