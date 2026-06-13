from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA_ROOT = ROOT / "output" / "time_step_coverage_audit" / "L300_350_400_ustar_grid_n4096_s3_plus_extra"
PACKAGE_DIR = DATA_ROOT / "result_package"
REPORT_DIR = ROOT / "output" / "reports" / "galloping_multispan_plus_extra"
DOCX_PATH = REPORT_DIR / "Galloping_Multispan_TimeDomain_Coverage_Report.docx"
CSV_PATH = PACKAGE_DIR / "coverage_big_table_source.csv"
RELIABILITY_CSV = PACKAGE_DIR / "geometry_reliability_summary.csv"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
TEXT = RGBColor(32, 32, 32)
HEADER_FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"
WARNING_FILL = "FFF2CC"
BORDER = "B7C9DD"


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def pct(value: str | float) -> str:
    return f"{100.0 * float(value):.2f}%"


def set_run_font(run, east_asia: str = "Microsoft YaHei") -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = "Galloping multi-span time-domain coverage study"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED
        set_run_font(run)

    footer = section.footer.paragraphs[0]
    footer.text = "Generated from pylon1 simulation outputs"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED
        set_run_font(run)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table, color: str = BORDER) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 9.0, fill: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    set_run_font(run)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        set_cell_shading(cell, fill)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run)


def add_para(doc: Document, text: str, *, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    for part in text.split("`"):
        run = p.add_run(part)
        set_run_font(run)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        set_run_font(run)


def add_callout(doc: Document, title: str, text: str, fill: str = BLUE_FILL) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "D9E2F3")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    set_run_font(r)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(text)
    r2.font.size = Pt(9.5)
    set_run_font(r2)


def add_simple_table(doc: Document, headers: list[str], rows: list[list[str]], *, font_size: float = 8.5) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, size=font_size, fill=HEADER_FILL)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            set_cell_text(cells[i], value, size=font_size)


def add_picture(doc: Document, path: Path, caption: str, width: float = 6.65) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = MUTED
    set_run_font(r)


def group_by_l(rows: list[dict[str, str]]) -> dict[float, list[dict[str, str]]]:
    grouped: dict[float, list[dict[str, str]]] = {}
    for row in rows:
        grouped.setdefault(float(row["L_m"]), []).append(row)
    for value in grouped:
        grouped[value].sort(key=lambda item: float(item["u_star"]))
    return grouped


def top_result_rows(rows: list[dict[str, str]]) -> list[list[str]]:
    selected = []
    for row in rows:
        if float(row["u_star"]) in {0.5, 0.6, 0.8, 1.0}:
            selected.append(
                [
                    f"{float(row['L_m']):.0f}",
                    f"{float(row['Sag_m']):.3f}",
                    f"{float(row['u_star']):.2f}",
                    f"{row['completed_n']}/{row['planned_n']}",
                    pct(row["failed_fraction"]),
                    pct(row["C6_large_response"]),
                    pct(row["C4_rolling_response_growth"]),
                    pct(row["C2_and_C4"]),
                ]
            )
    return selected


def build_report() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    rows = load_csv(CSV_PATH)
    reliability = load_csv(RELIABILITY_CSV)

    doc = Document()
    style_doc(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Multi-Span Time-Domain Galloping Coverage Report")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = DARK_BLUE
    set_run_font(r)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("基于 L-u_star 参数化 OpenSees 时程模拟的覆盖率判据研究")
    r.font.size = Pt(12)
    r.font.color.rgb = MUTED
    set_run_font(r)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(f"Updated: {date.today().isoformat()} | Dataset: L300/350/400 plus supplementary seeds")
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    set_run_font(r)

    add_callout(
        doc,
        "Executive finding / 核心结论",
        "在当前 Zebra ACSR、parabolic_tension 下垂匹配和 4096 步时域窗口中，C6 大响应覆盖率随 u_star 增强最清晰，"
        "是当前阶段最可靠的 developed-response 曲面。C4 反映增长过程但受随机风场 realization 影响；C2 是有效负阻尼机制指标，"
        "不宜单独解释为严重程度曲线；C2&C4 是保守同步判据。高失败率格点应作为非正常动力终止/数值失稳辅助曲面保留。"
    )

    add_heading(doc, "1. 研究框架与方法论", 1)
    add_para(
        doc,
        "本研究的目标不是只寻找单一临界风速，而是构建 `coverage(L, u_star)` 型二维响应函数。"
        "其中 `L` 是跨距，`u_star` 是摩阻风速输入，覆盖率定义为在 204.8 s 时程中满足某一 galloping 判据的时间步比例。"
        "这种定义把 galloping 从二值事件扩展为时域概率/占有率指标，更适合比较不同判据的覆盖范围、灵敏度和工程解释。"
    )
    add_para(
        doc,
        "当前工作流由四层组成：几何层、风荷载层、有限元时程层和判据后处理层。几何层采用 Zebra ACSR 基线参数，"
        "跨距取 300、350、400 m；每个跨距的下垂高度由同一导线自重和初拉力的浅弧垂公式匹配。风荷载层使用从 MATLAB 原始工程移植的"
        " friction-based generator，通过 `U(z)=u_star/0.387*log(z/0.05)` 生成平均风，并叠加 Kaimal/Davenport 型随机湍流与相干性。"
        "有限元层调用 OpenSees 时程分析，记录位移、加速度、反力和单元有效阻尼。后处理层把所有判据落到时间步，并计算覆盖率。"
    )
    add_numbered(
        doc,
        [
            "几何匹配：使用 `H(L)=wL^2/(8T0)`，其中 `w=15.90 N/m`，`T0=19.785 kN`。得到 H(300)=9.041 m，H(350)=12.306 m，H(400)=16.073 m。",
            "风场输入：把 `u_star` 作为主要自变量。当前范围为 0.30、0.40、0.50、0.60、0.80、1.00 m/s。",
            "随机性处理：每个格点至少 3 个 seed；高失败或 completed 样本不足的格点补充 2 个 seed。",
            "固定审计窗口：每个完成样本审计前 4096 条目标记录，避免自适应子步导致高风速样本在覆盖率计算中被过度加权。",
            "失败样本处理：非正常 OpenSees 终止不进入完成样本平均，但作为 `failed_fraction(L,u_star)` 单独报告。"
        ],
    )

    add_heading(doc, "2. 标准设计与判据解释", 1)
    add_para(
        doc,
        "判据设计遵循一个基本原则：物理机制、响应表现和工程后果必须分开记录。Den Hartog/effective damping 机制能够解释自激来源，"
        "但不能保证已经出现大幅响应；响应增长判据能够观察时域放大，但可能受随机风和窗口位置影响；工程限值判据对应后果，"
        "但不等同于物理 onset。因此，本阶段不把判据强行取交集，而是分别形成覆盖率曲线或曲面。"
    )
    add_simple_table(
        doc,
        ["判据", "定义角色", "当前实现", "报告解释"],
        [
            ["C2 负阻尼", "机制指标", "足够比例单元在同一时间步出现负有效阻尼", "表示自激能量输入条件；不是单调严重程度曲线"],
            ["C4 增长", "响应过程指标", "20 s rolling response p95 相对前窗/基线增长", "用于识别 galloping-like growth；对 seed 和窗口敏感"],
            ["C6 大响应", "发展响应指标", "位移超过 sag-based 大响应阈值", "当前最清晰的严重响应曲面，适合作为主结果之一"],
            ["C2&C4", "保守同步指标", "同一时间步同时满足 C2 与 C4", "严格但覆盖率低，可作为高置信机制-响应交集"],
            ["C7", "工程后果指标", "当前 clearance proxy", "尚未按真实净空限值校准，不用于 primary galloping 结论"],
            ["失败率", "辅助失稳指标", "完成前非正常终止 / planned cases", "代表强非线性响应或数值失稳风险，不应直接丢弃"],
        ],
        font_size=8.0,
    )
    add_callout(
        doc,
        "C7 使用限制",
        "当前 C7 覆盖率接近 99%，说明 clearance proxy 设置过于宽泛或与当前几何基准不完全对应。"
        "因此 C7 只作为工程限值通道保留，后续必须用真实相间距、地线/地面净空、tower geometry 或 utility limit 重新校准。",
        fill=WARNING_FILL,
    )

    add_heading(doc, "3. 实验矩阵与数据完整性", 1)
    add_para(
        doc,
        "最新数据集包含原始 54 个 multi-span 工况与 10 个补充工况，共 64 个 planned cases。"
        "补样集中在失败率高或完成样本偏少的格点：L=300,u_star=0.80；L=350,u_star=0.80 和 1.00；L=400,u_star=0.50 和 0.80。"
        "最终完成样本 48 个，非正常终止样本 16 个。"
    )
    add_simple_table(
        doc,
        ["L (m)", "H (m)", "u_star", "completed/planned", "failed fraction"],
        [
            [
                f"{float(row['L_m']):.0f}",
                f"{float(row['Sag_m']):.3f}",
                f"{float(row['u_star']):.2f}",
                f"{row['completed_n']}/{row['planned_n']}",
                pct(row["failed_fraction"]),
            ]
            for row in reliability
        ],
        font_size=7.6,
    )

    add_heading(doc, "4. 覆盖率结果", 1)
    add_para(
        doc,
        "下表摘录中高风速区间的关键结果。完整分线形大表已另存为 Excel，并在附图中按 L-H 构型分表展示。"
    )
    add_simple_table(
        doc,
        ["L", "H", "u_star", "completed/planned", "failed", "C6", "C4", "C2&C4"],
        top_result_rows(rows),
        font_size=7.6,
    )
    add_picture(
        doc,
        PACKAGE_DIR / "coverage_tables_all_geometries.png",
        "Figure 1. Coverage table by geometry. Main table columns follow the requested order: u_star, C2&C4, C2 negative damping, C4 growth, C6 large response, and C7.",
        width=6.85,
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "5. 二维曲面结果", 1)
    add_para(
        doc,
        "每个判据单独绘制为 L-u_star 二维曲面，避免把机制、响应和后果混成一个总分。"
        "图中颜色和数字均为覆盖率百分比；失败率曲面表示 planned cases 中非正常终止的比例。"
    )
    for filename, caption in [
        ("C6_large_response_surface.png", "Figure 2. C6 developed large-response coverage surface."),
        ("C4_growth_surface.png", "Figure 3. C4 response-growth coverage surface."),
        ("C2_negative_damping_surface.png", "Figure 4. C2 effective negative-damping coverage surface."),
        ("C2C4_surface.png", "Figure 5. Conservative simultaneous C2&C4 coverage surface."),
        ("failed_fraction_surface.png", "Figure 6. Failed/non-normal termination fraction surface."),
        ("C7_surface.png", "Figure 7. Current C7 clearance-proxy coverage surface; retained for tracking but not calibrated for primary conclusions."),
    ]:
        add_picture(doc, PACKAGE_DIR / filename, caption, width=5.95)

    add_heading(doc, "6. 结果解释", 1)
    add_bullets(
        doc,
        [
            "C6 是当前最稳定的 developed-response 指标。对三个跨距，C6 覆盖率随 u_star 从 0.50-0.60 开始快速升高，并在 0.80-1.00 区间达到高覆盖率。",
            "C4 的覆盖率不严格单调，说明响应增长受随机风 realization、非线性平台和窗口位置影响。它仍然重要，因为它描述了从机制到响应的增长过程。",
            "C2 在低风速下也有稳定占有率，但随大响应发展反而下降。这并不表示风险降低，而是说明导线进入大幅运动后，瞬时角攻角/阻尼状态不再长期停留在小扰动负阻尼区。",
            "C2&C4 覆盖率很低，符合其保守同步判据定位。它适合识别机制和响应在同一时间步重合的高置信窗口，不适合作为唯一 galloping 指标。",
            "失败率在 L=300,u=0.80、L=350,u=0.80/1.00、L=400,u=0.50/0.80 仍然较高。结合输出文件审计，这些失败更接近强非线性动力响应下的非正常终止，而不是前处理错误。",
        ],
    )

    add_heading(doc, "7. 局限性与下一步", 1)
    add_bullets(
        doc,
        [
            "当前 L 只有 300、350、400 m 三档，足以展示趋势，但不足以拟合最终连续二维函数。建议下一轮加入 250、322.8、450、500 m。",
            "C7 必须重新校准。当前 clearance proxy 过高覆盖，不能用于工程失效概率或 primary galloping conclusion。",
            "失败率应与完成样本覆盖率并列展示。后续可以对失败前窗口单独截取，计算 termination-before-failure 的 C2/C4/C6 状态。",
            "若要把覆盖率转为临界边界，可选 iso-coverage 曲线，例如 C6=50%、C6=80%、failed_fraction=20%。但当前阶段更适合先建立趋势曲面。",
            "束导线、覆冰截面和扭转耦合尚未进入主模型。Den Hartog/C2 当前主要对应垂向 galloping，Nigol/Clarke 型扭转机制应作为后续扩展。"
        ],
    )

    add_heading(doc, "8. 文件索引", 1)
    add_simple_table(
        doc,
        ["Artifact", "Path"],
        [
            ["Result package", str(PACKAGE_DIR.relative_to(ROOT))],
            ["Source table", str((PACKAGE_DIR / "coverage_big_table_source.csv").relative_to(ROOT))],
            ["Excel big table", str((PACKAGE_DIR / "coverage_big_table_by_geometry.xlsx").relative_to(ROOT))],
            ["Reliability summary", str((PACKAGE_DIR / "geometry_reliability_summary.csv").relative_to(ROOT))],
            ["Time-step audits", str((DATA_ROOT / "cases").relative_to(ROOT))],
        ],
        font_size=8.0,
    )

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_report()
